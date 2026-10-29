"""
generar_reporte_excel.py (GENERICO, copiar tal cual a cada proyecto -
ver NUEVO_PROYECTO_CHECKLIST.md, no hace falta editarlo por proyecto).

Genera un Excel consolidado ("Resumen" por RF + "Detalle por caso") del
ESTADO MAS RECIENTE conocido de cada CP, leyendo:
  - matriz_casos.json (para RF/CU/Nombre de cada caso)
  - todos los reportes/reporte_*.csv que dejo run_all.py (toma, para cada
    CP, el resultado del CSV mas nuevo que lo menciona - asi que corridas
    parciales despues de una corrida completa SI actualizan el estado)
  - casos_fallidos_confirmados.json, SI EXISTE en la carpeta del proyecto
    (opcional - ver metodologia abajo)

No ejecuta ningun caso de prueba - solo lee lo que ya se corrio. Si un CP
nunca aparecio en ningun CSV, queda como "SIN CORRER".

**Metodologia de 3 estados (agregada PR000335-2, 15/09/2026):** un caso
que no dio OK puede ser dos cosas MUY distintas - (a) la automatizacion
todavia no esta bien (selector viejo, falta un dato de prueba, timing) o
(b) se confirmo EN VIVO que es un defecto REAL de la app, listo para
reportar a Desarrollo. Mezclar ambas en un solo "Fallido" rojo hace que
el reporte no sirva para decidir que reportar todavia. Si existe
`casos_fallidos_confirmados.json` junto a este script (formato:
`{"CPxxx": {"motivo": "...", "fecha_verificacion": "...", "evidencia":
"..."}, ...}` - ver PR000335-2 como referencia, incluye tambien como
reproducirlo con `sesion_manual_con_red.py`), los CP que estan ahi se
muestran como "Fallido" (rojo) y el resto de los que no dieron OK se
muestran como "NO OK" (amarillo, pendiente de revisar/confirmar) en vez
de mezclarse. Sin ese archivo (default), TODO lo que no sea OK cae en
"NO OK" - un proyecto que no lo use no pierde nada, solo no tiene esa
distincion todavia.

Uso:
    .venv\\Scripts\\python.exe generar_reporte_excel.py
"""
import csv
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
MATRIZ = json.loads((BASE / "matriz_casos.json").read_text(encoding="utf-8"))
REPORTES_DIR = BASE / "reportes"
CASOS_FALLIDOS_PATH = BASE / "casos_fallidos_confirmados.json"

VERDE = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ROJO = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
AMARILLO = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
GRIS_CLARO = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
AZUL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
GRIS_RF = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
FUENTE_HEADER = Font(color="FFFFFF", bold=True)
FUENTE_RF = Font(bold=True)
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

_PATRON_TIMESTAMP = re.compile(r"reporte_(\d{8}_\d{6})\.csv$")


def cargar_casos_fallidos_confirmados():
    """CP -> dict(motivo, fecha_verificacion, evidencia) para los casos
    verificados EN VIVO como defecto real (no bug de automatizacion).
    Devuelve {} si el archivo no existe en este proyecto (opcional)."""
    if not CASOS_FALLIDOS_PATH.exists():
        return {}
    datos = json.loads(CASOS_FALLIDOS_PATH.read_text(encoding="utf-8"))
    return {k: v for k, v in datos.items() if not k.startswith("_")}


def _normalizar_resultado(texto, cp, confirmados):
    t = (texto or "").strip().upper()
    if t == "OK":
        return "OK"
    if t in ("FALLO", "FALLIDO"):
        return "FALLIDO" if cp in confirmados else "NO OK"
    if t == "TIMEOUT":
        return "TIMEOUT"
    return t or "?"


def recolectar_resultados(confirmados):
    """Recorre todos los reporte_*.csv en orden cronologico (por el
    timestamp del nombre de archivo) y se queda con el ULTIMO resultado
    visto de cada CP - asi una corrida parcial mas reciente pisa a una
    corrida completa mas vieja para esos CP puntuales."""
    csvs = sorted(
        REPORTES_DIR.glob("reporte_*.csv"),
        key=lambda p: _PATRON_TIMESTAMP.search(p.name).group(1) if _PATRON_TIMESTAMP.search(p.name) else "",
    )
    resultados = {}  # cp -> (resultado, detalle, archivo_origen, timestamp)
    for ruta in csvs:
        m = _PATRON_TIMESTAMP.search(ruta.name)
        ts = m.group(1) if m else ""
        try:
            with open(ruta, encoding="utf-8-sig", newline="") as f:
                for fila in csv.DictReader(f, delimiter=";"):
                    cp = (fila.get("CP") or "").strip()
                    if not cp:
                        continue
                    resultados[cp] = (
                        _normalizar_resultado(fila.get("Resultado"), cp, confirmados),
                        (fila.get("Detalle") or "").strip(),
                        ruta.name,
                        ts,
                    )
        except Exception as e:
            print(f"  (aviso) no se pudo leer {ruta.name}: {e}")
    return resultados


def construir_datos(resultados, confirmados):
    """Calcula el resumen por RF + detalle por caso a partir de
    `resultados` (ver recolectar_resultados()) y `confirmados` (ver
    cargar_casos_fallidos_confirmados()), en un dict plano JSON-friendly
    que tanto construir_workbook() como el webapp (server.py) consumen
    sin recalcular nada por su cuenta."""
    conteos = {}
    for cp, v in MATRIZ.items():
        rf = v["rf"]
        estado = resultados.get(cp, ("SIN CORRER", "", "", ""))[0]
        conteos.setdefault(rf, {"OK": 0, "NO OK": 0, "FALLIDO": 0, "SIN CORRER": 0, "OTRO": 0})
        clave = estado if estado in ("OK", "NO OK", "FALLIDO", "SIN CORRER") else "OTRO"
        conteos[rf][clave] += 1

    resumen = []
    total = {"OK": 0, "NO OK": 0, "FALLIDO": 0, "SIN CORRER": 0, "OTRO": 0}
    for rf in sorted(conteos):
        c = conteos[rf]
        n_total = sum(c.values())
        pct = round(c["OK"] / n_total * 100) if n_total else None
        resumen.append({
            "rf": rf, "ok": c["OK"], "no_ok": c["NO OK"] + c["OTRO"],
            "fallido": c["FALLIDO"], "sin_correr": c["SIN CORRER"],
            "total": n_total, "pct_ok": pct,
        })
        for k in total:
            total[k] += c[k]

    n_general = sum(total.values())
    pct_general = round(total["OK"] / n_general * 100) if n_general else None
    total_row = {
        "ok": total["OK"], "no_ok": total["NO OK"] + total["OTRO"],
        "fallido": total["FALLIDO"], "sin_correr": total["SIN CORRER"],
        "total": n_general, "pct_ok": pct_general,
    }

    detalle = []
    for cp in sorted(MATRIZ, key=lambda k: int(re.sub(r"\D", "", k) or 0)):
        v = MATRIZ[cp]
        estado, detalle_txt, archivo, ts = resultados.get(cp, ("SIN CORRER", "", "", ""))
        ts_legible = f"{ts[:4]}-{ts[4:6]}-{ts[6:8]} {ts[9:11]}:{ts[11:13]}" if len(ts) >= 13 else ""
        motivo = confirmados.get(cp, {}).get("motivo", "") if estado == "FALLIDO" else ""
        detalle.append({
            "cp": cp, "rf": v["rf"], "cu": v["cu"], "nombre": v["name"],
            "resultado": estado, "ultima_corrida": ts_legible,
            "detalle": detalle_txt, "motivo_confirmado": motivo,
        })

    sin_correr = [cp for cp in MATRIZ if cp not in resultados]
    return {"resumen": resumen, "total": total_row, "detalle": detalle, "sin_correr": sin_correr}


def construir_workbook(datos):
    wb = Workbook()

    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append(["Reporte consolidado de ejecución", "", "", "", "", ""])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append([f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    resumen.append([])
    resumen.append(["RF", "OK", "NO OK", "Fallido", "Sin correr", "Total", "% OK"])
    for c in resumen[4]:
        c.font = FUENTE_HEADER
        c.fill = AZUL_HEADER

    for fila_rf in datos["resumen"]:
        pct = f"{fila_rf['pct_ok']}%" if fila_rf["pct_ok"] is not None else "-"
        resumen.append([fila_rf["rf"], fila_rf["ok"], fila_rf["no_ok"], fila_rf["fallido"], fila_rf["sin_correr"], fila_rf["total"], pct])

    t = datos["total"]
    pct_general = f"{t['pct_ok']}%" if t["pct_ok"] is not None else "-"
    resumen.append(["TOTAL", t["ok"], t["no_ok"], t["fallido"], t["sin_correr"], t["total"], pct_general])
    for c in resumen[resumen.max_row]:
        c.font = Font(bold=True)
        c.fill = GRIS_RF

    for i, a in enumerate([10, 8, 8, 10, 12, 8, 8], start=1):
        resumen.column_dimensions[get_column_letter(i)].width = a

    detalle = wb.create_sheet("Detalle por caso")
    detalle.append(["CP", "RF", "CU", "Nombre del caso", "Resultado", "Última corrida", "Detalle", "Motivo confirmado (si Fallido)"])
    for c in detalle[1]:
        c.font = FUENTE_HEADER
        c.fill = AZUL_HEADER

    fila = 2
    rf_actual = None
    fill_por_estado = {"OK": VERDE, "FALLIDO": ROJO, "NO OK": AMARILLO, "SIN CORRER": GRIS_CLARO}
    for d in datos["detalle"]:
        if d["rf"] != rf_actual:
            rf_actual = d["rf"]
            detalle.append([rf_actual, "", "", "", "", "", "", ""])
            detalle.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=8)
            for c in detalle[fila]:
                c.fill = GRIS_RF
                c.font = FUENTE_RF
            fila += 1

        detalle.append([d["cp"], d["rf"], d["cu"], d["nombre"], d["resultado"], d["ultima_corrida"], d["detalle"], d["motivo_confirmado"]])
        celda_estado = detalle.cell(row=fila, column=5)
        celda_estado.fill = fill_por_estado.get(d["resultado"], AMARILLO)
        celda_estado.font = Font(bold=True)
        celda_estado.alignment = Alignment(horizontal="center")
        for col in range(1, 9):
            detalle.cell(row=fila, column=col).border = BORDE
            detalle.cell(row=fila, column=col).alignment = Alignment(vertical="top", wrap_text=(col in (4, 7, 8)))
        fila += 1

    for i, a in enumerate([8, 8, 8, 42, 12, 16, 60, 60], start=1):
        detalle.column_dimensions[get_column_letter(i)].width = a
    detalle.freeze_panes = "A2"

    return wb


def main():
    confirmados = cargar_casos_fallidos_confirmados()
    resultados = recolectar_resultados(confirmados)
    datos = construir_datos(resultados, confirmados)
    wb = construir_workbook(datos)

    REPORTES_DIR.mkdir(exist_ok=True)
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    destino_xlsx = REPORTES_DIR / f"reporte_consolidado_{marca}.xlsx"
    destino_json = REPORTES_DIR / f"reporte_consolidado_{marca}.json"
    wb.save(destino_xlsx)
    destino_json.write_text(
        json.dumps({"generado": datetime.now().isoformat(), **datos}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"Reporte generado: {destino_xlsx}")
    print(f"Reporte JSON generado: {destino_json}")
    if confirmados:
        print(f"  ({len(confirmados)} caso(s) marcados como Fallido confirmado via casos_fallidos_confirmados.json)")
    if datos["sin_correr"]:
        print(f"  ({len(datos['sin_correr'])} caso(s) nunca aparecieron en ningun CSV: {', '.join(datos['sin_correr'])})")


if __name__ == "__main__":
    main()
